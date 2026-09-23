# -*- coding: utf-8 -*-
"""
Genera la propuesta comercial (Word .docx) para Gael - Metales Infinity.
Marca: Aimarktech. Paleta y tipografias tomadas de soyaimarktech.com.
v2: agrega Kit de Marca, auditoria de redes, calendario por etapas y beneficios.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ------------------------------------------------------------------ paleta
BLUE = "0A74DA"; CYAN = "00C2FF"; YELLOW = "FFCE00"; GREEN = "28A745"
INK = "0E1B2C"; INK_SOFT = "46566B"; BG_SOFT = "F4F8FD"; DARK = "071426"
LINE = "E3EBF4"; WHITE = "FFFFFF"; RED = "C62F3B"
HEAD = "Montserrat"; BODY = "Lato"

# ------------------------------------------------------------------ helpers
def _set_font(run, font=BODY, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = font; run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), font)

def para(host, text="", size=10.5, color=INK, bold=False, italic=False,
         font=BODY, align=None, space_after=6, space_before=0, line=1.15):
    p = host.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before); pf.line_spacing = line
    if text:
        r = p.add_run(text); _set_font(r, font, size, color, bold, italic)
    return p

def runs(p, parts):
    for text, kw in parts:
        r = p.add_run(text); _set_font(r, **kw)
    return p

def bottom_border(p, color=BLUE, sz=14):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(sz)); b.set(qn('w:space'), '4'); b.set(qn('w:color'), color)
    pbdr.append(b); pPr.append(pbdr)

def cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right), ('left', left), ('right', right)):
        e = OxmlElement('w:' + tag); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)

def grid_borders(table, color=LINE, sz=6):
    tblPr = table._tbl.tblPr; borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color); borders.append(e)
    tblPr.append(borders)

def top_accent(cell, color=BLUE, sz=24):
    tcPr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    t = OxmlElement('w:top'); t.set(qn('w:val'), 'single'); t.set(qn('w:sz'), str(sz))
    t.set(qn('w:space'), '0'); t.set(qn('w:color'), color); borders.append(t); tcPr.append(borders)

def section_title(doc, number, text):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(15); pf.space_after = Pt(7); pf.keep_with_next = True
    r = p.add_run(f"{number}  "); _set_font(r, HEAD, 13.5, CYAN, bold=True)
    r = p.add_run(text); _set_font(r, HEAD, 13.5, BLUE, bold=True)
    bottom_border(p, LINE, 10)
    return p

def bullet(doc, label, text, dot="●", dotcolor=BLUE):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.line_spacing = 1.12
    r = p.add_run(dot + " "); _set_font(r, BODY, 10.5, dotcolor, bold=True)
    if label:
        r = p.add_run(label + " "); _set_font(r, BODY, 10.5, INK, bold=True)
    r = p.add_run(text); _set_font(r, BODY, 10.5, INK_SOFT)
    return p

def callout(doc, parts, accent=BLUE, fill=BG_SOFT):
    box = doc.add_table(rows=1, cols=1); box.columns[0].width = CONTENT_W
    c = box.cell(0, 0); c.width = CONTENT_W
    cell_bg(c, fill); cell_margins(c, 160, 160, 220, 220); top_accent(c, accent, 20)
    c.paragraphs[0].text = ""; runs(c.paragraphs[0], parts)
    return box

def simple_table(doc, headers, rows, widths):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid_borders(t, LINE, 6)
    ws = [int(CONTENT_W * w) for w in widths]
    hc = t.rows[0].cells
    for j, htxt in enumerate(headers):
        hc[j].width = ws[j]; cell_bg(hc[j], BLUE); cell_margins(hc[j])
        hc[j].paragraphs[0].text = ""; r = hc[j].paragraphs[0].add_run(htxt); _set_font(r, HEAD, 9.5, WHITE, bold=True)
    for i, row in enumerate(rows, start=1):
        fill = WHITE if i % 2 else BG_SOFT
        for j, val in enumerate(row):
            cc = t.rows[i].cells[j]; cc.width = ws[j]; cell_bg(cc, fill); cell_margins(cc)
            cc.paragraphs[0].text = ""
            r = cc.paragraphs[0].add_run(val); _set_font(r, BODY, 9.7, INK if j == 0 else INK_SOFT, bold=(j == 0))
    return t

# ------------------------------------------------------------------ documento
doc = Document()
st = doc.styles['Normal']; st.font.name = BODY; st.font.size = Pt(10.5)
st.font.color.rgb = RGBColor.from_string(INK)
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
CONTENT_W = sec.page_width - sec.left_margin - sec.right_margin

foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = foot.add_run("Aimarktech  ·  Propuesta para Metales Infinity  ·  Documento confidencial")
_set_font(r, BODY, 8, INK_SOFT)

# ============================================================== PORTADA
cover = doc.add_table(rows=1, cols=1); cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.columns[0].width = CONTENT_W
c = cover.cell(0, 0); c.width = CONTENT_W
cell_bg(c, DARK); cell_margins(c, 360, 360, 360, 360); c.paragraphs[0].text = ""
runs(c.paragraphs[0], [("AIMARKTECH", dict(font=HEAD, size=13, color=CYAN, bold=True)),
                       ("    IA · Marketing · Mentalidad empresarial", dict(size=9.5, color="B9C6D6"))])
para(c, "Propuesta: tu marca y tu tienda en línea", size=24, color=WHITE, bold=True, font=HEAD, space_before=10, space_after=2)
para(c, "Metales Infinity", size=24, color=YELLOW, bold=True, font=HEAD, space_after=10)
para(c, "Una tienda propia bilingüe (español / inglés) para vender con orden, cobrar con "
        "anticipo y depender menos del día a día.", size=11.5, color="D7E4F0", space_after=4)

meta = doc.add_table(rows=1, cols=1); meta.columns[0].width = CONTENT_W
mc = meta.cell(0, 0); mc.width = CONTENT_W
cell_bg(mc, BG_SOFT); cell_margins(mc, 200, 200, 240, 240); top_accent(mc, BLUE, 24); mc.paragraphs[0].text = ""
runs(mc.paragraphs[0], [("Preparado para:  ", dict(bold=True)), ("Gael — Metales Infinity · San Juan del Río, Qro.", dict(color=INK_SOFT))])
runs(para(mc, space_after=3), [("Preparado por:  ", dict(bold=True)), ("José Antonio Aguilar · Aimarktech", dict(color=INK_SOFT))])
runs(para(mc, space_after=0), [("Fecha:  ", dict(bold=True)), ("__________________     ·     Septiembre 2026", dict(color=INK_SOFT))])

# ------ Antes de empezar
callout(doc, [
    ("Antes de empezar.  ", dict(bold=True, color=BLUE, size=11, font=HEAD)),
    ("Gael, esto no es una lista de servicios ni promesas de \u201cposts bonitos\u201d. Es el "
     "resultado de mirar tu negocio con tus propios números. Tu negocio no está mal: tienes "
     "lo más difícil —buen producto, clientes que regresan y calidad que te distingue—. Lo "
     "que falta es un sistema para que no todo dependa de ti. Eso es lo que te proponemos, "
     "con honestidad y sin ofrecerte lo que no se pueda cumplir.", dict(color=INK_SOFT)),
])

# ============================================================== 1. EN CORTO
section_title(doc, "1", "En corto")
para(doc, "Metales Infinity fabrica sellos y figuras vaqueras personalizadas de calidad, y hoy "
          "vende principalmente por Facebook y WhatsApp. El siguiente paso es tu propia tienda "
          "en línea, bilingüe, donde tus clientes vean productos y precios, puedan pedir y "
          "—cuando lo decidas— pagar en línea con un anticipo.", color=INK_SOFT)
para(doc, "No te prometemos un número mágico de ventas. Buscamos orden: menos tiempo "
          "contestando lo mismo, menos riesgo de cancelaciones y una vitrina profesional que "
          "trabaje por ti las 24 horas, también para el mercado latino de Estados Unidos.", color=INK_SOFT)

# ============================================================== 2. RADIOGRAFÍA
section_title(doc, "2", "Radiografía de tu negocio")
para(doc, "Esto es lo que dicen los números que compartiste (los confirmamos y afinamos en la "
          "Etapa 1):", color=INK_SOFT, space_after=8)
simple_table(doc,
    ["Indicador", "Hoy (según tus datos)", "Qué significa"],
    [["Capacidad usada", "40 de 120 piezas/día (~33%)", "Capacidad instalada para crecer sin comprar maquinaria"],
     ["Tiempo en atención", "~6 h/día en mensajes", "Hoy el negocio depende de ti"],
     ["Pedidos por mes", "de ~8 (bajo) a 60 (alto)", "Mucha variación; conviene estabilizar"],
     ["Ticket común", "~$3,000 por juego de 20", "Márgenes sanos, con espacio para volumen"],
     ["Recompra", "pines ~15 días · sellos ~30–45 días", "Clientes que regresan, hoy sin seguimiento"],
     ["Mercado", "~80% México / ~20% EE.UU.", "El mercado latino de EE.UU. ya te compra"]],
    [0.22, 0.30, 0.48])
para(doc, "", space_after=2)
callout(doc, [("La foto real.  ", dict(bold=True, color=INK)),
    ("Metales Infinity es un negocio rentable que hoy descansa por completo en ti: eres el "
     "único motor de ventas, atención y logística. El objetivo es darle un sistema que venda, "
     "cobre y dé seguimiento, para que no todo pase por tus manos.", dict(color=INK_SOFT))], accent=BLUE)
para(doc, "En concreto, cuatro cosas te frenan hoy:", color=INK_SOFT, space_before=6, space_after=6)
bullet(doc, "Todo depende de ti.", "Muchas horas al día contestando y mandando fotos una por una; si faltas, la venta se detiene.")
bullet(doc, "Sin catálogo con precios visibles.", "Cada venta es manual y lenta; el cliente no puede explorar a su ritmo.")
bullet(doc, "Sin anticipo.", "Fabricas primero y cobras al final: si cancelan, tú absorbes material y tiempo.")
bullet(doc, "Sin tienda propia.", "Competidores como Sellare ya venden con tienda y precios visibles, y te toman ventaja en lo digital. En Querétaro casi no tienes competencia directa: la ventana está abierta.")
callout(doc, [("Sobre lo fiscal (opcional y a tu ritmo).  ", dict(bold=True, color=INK)),
    ("Darte de alta en el SAT no es necesario para arrancar tu tienda en México ni para tener "
     "el sitio en inglés. Solo se vuelve útil más adelante si quieres facturar formalmente o "
     "cobrar y retirar en dólares. Lo dejamos como referencia, tú decides cuándo.", dict(color=INK_SOFT))], accent=YELLOW)

# ============================================================== 3. PROPUESTA
section_title(doc, "3", "La propuesta: tu propia tienda en línea")
para(doc, "Una tienda propia (no rentada): el sitio es tuyo, en tu dominio, con el código "
          "respaldado y alojado en la nube. Bilingüe desde el inicio, con precios visibles, "
          "catálogo, pedidos personalizados, reseñas de clientes y botón de WhatsApp. Para "
          "arrancar, dos caminos:", color=INK_SOFT, space_after=10)

opts = doc.add_table(rows=1, cols=2); opts.alignment = WD_TABLE_ALIGNMENT.CENTER
half = int(CONTENT_W / 2)
for i in range(2):
    opts.columns[i].width = half; opts.cell(0, i).width = half
grid_borders(opts, WHITE, 8)
o1 = opts.cell(0, 0); cell_bg(o1, BG_SOFT); cell_margins(o1, 200, 200, 220, 220); top_accent(o1, INK_SOFT, 20); o1.paragraphs[0].text = ""
para(o1, "OPCIÓN 1", size=9, color=INK_SOFT, bold=True, font=HEAD, space_after=1)
para(o1, "Catálogo + compra por WhatsApp", size=12, color=INK, bold=True, font=HEAD, space_after=2)
para(o1, "Ligera y rápida", size=9.5, color=INK_SOFT, italic=True, space_after=6)
para(o1, "Tu catálogo con fotos, medidas y precios; el cliente pide con un botón que abre "
         "WhatsApp con el mensaje ya escrito.", size=10, color=INK_SOFT, space_after=6)
runs(para(o1, space_after=3), [("A favor:  ", dict(bold=True, size=10, color=GREEN)), ("rápida y económica; conservas tu trato personal.", dict(size=10, color=INK_SOFT))])
runs(para(o1, space_after=0), [("Límite:  ", dict(bold=True, size=10, color=RED)), ("no cobra en línea; no es autoservicio completo.", dict(size=10, color=INK_SOFT))])
o2 = opts.cell(0, 1); cell_bg(o2, BG_SOFT); cell_margins(o2, 200, 200, 220, 220); top_accent(o2, BLUE, 26); o2.paragraphs[0].text = ""
runs(o2.paragraphs[0], [("OPCIÓN 2   ", dict(size=9, color=BLUE, bold=True, font=HEAD)), ("\u2605 RECOMENDADA", dict(size=8.5, color=YELLOW, bold=True, font=HEAD))])
para(o2, "Tienda con carrito y pago en línea", size=12, color=INK, bold=True, font=HEAD, space_after=2)
para(o2, "Completa", size=9.5, color=INK_SOFT, italic=True, space_after=6)
para(o2, "Además del catálogo: carrito, pago en línea (tarjeta / Mercado Pago), anticipo y un "
         "flujo especial para personalizados (cotización → anticipo → aprobación → fabricación "
         "→ liquidación). Incluye un panel para que tú administres productos, precios y pedidos.",
     size=10, color=INK_SOFT, space_after=6)
runs(para(o2, space_after=3), [("A favor:  ", dict(bold=True, size=10, color=GREEN)), ("vende 24/7, cobra anticipo, ordena pedidos y saldos, lista para EE.UU.", dict(size=10, color=INK_SOFT))])
runs(para(o2, space_after=0), [("Necesita:  ", dict(bold=True, size=10, color=RED)), ("tener claros tus precios y el medio de pago.", dict(size=10, color=INK_SOFT))])
para(doc, "", space_after=2)
callout(doc, [("Nuestra recomendación.  ", dict(bold=True, color=BLUE)),
    ("Avanzar hacia la Opción 2 en su versión propia y bilingüe. Podemos lanzar rápido una "
     "primera versión (catálogo + compra) e ir sumando el pago en línea y el panel conforme lo "
     "validemos con tus clientes.", dict(color=INK_SOFT))], fill="EAF2FC")

# ============================================================== 4. KIT DE MARCA
section_title(doc, "4", "Punto de partida: tu Kit de Marca")
para(doc, "Una tienda profesional necesita una marca lista para web. Vimos que ya tienes un "
          "logotipo y una marca de agua en tus fotos; el primer paso es confirmar si existen en "
          "el formato y los tamaños que un sitio requiere. Si no, empezamos por ahí (Etapa 0).",
     color=INK_SOFT, space_after=6)
bullet(doc, "Logotipo en varios formatos y tamaños:", "versión horizontal, ícono/avatar, marca de agua y favicon; idealmente en vector (SVG/PDF) para que se vea nítido en cualquier medida.")
bullet(doc, "Identidad visual:", "paleta de colores oficial, tipografías y estilo de fotografía, para que todo se vea consistente.")
bullet(doc, "Guía breve de marca:", "cómo y dónde usar el logo, colores y mensajes clave (tu \u201ccalidad, precisión y esencia vaquera\u201d).")
callout(doc, [("Si aún no tienes todo esto,  ", dict(bold=True, color=INK)),
    ("no hay problema: lo armamos como Etapa 0 antes de construir el sitio, para no rehacer "
     "trabajo después.", dict(color=INK_SOFT))], accent=GREEN)

# ============================================================== 5. REDES
section_title(doc, "5", "Tus redes: una sola voz de marca")
para(doc, "Hoy tu fuerte es Facebook. Antes y durante el lanzamiento conviene revisar y unificar "
          "todas tus presencias para que el cliente reconozca la misma marca en todos lados y "
          "todas lleven a tu tienda:", color=INK_SOFT, space_after=6)
bullet(doc, "Auditoría:", "confirmar qué tienes activo (Facebook, Instagram, TikTok, Google/Maps) y qué falta abrir.")
bullet(doc, "Homogenizar perfiles:", "mismo nombre, foto, portada, biografía y enlaces en todas las redes, apuntando a la tienda.")
bullet(doc, "Google (opinión de clientes):", "hoy tienes pocas reseñas; montamos un proceso simple para pedir opiniones reales con foto.")
bullet(doc, "Contenido:", "aprovechar el proceso de fabricación (metal, corte, resultado en la tejana) que funciona muy bien en Reels/TikTok.")

# ============================================================== 6. PLAN Y CALENDARIO
section_title(doc, "6", "Plan de trabajo y calendario")
para(doc, "Trabajamos por etapas con una duración estimada. Las semanas se cuentan desde la "
          "fecha de arranque; al fijarla, cada etapa toma sus fechas exactas.", color=INK_SOFT, space_after=8)
simple_table(doc,
    ["Etapa", "Enfoque", "Entregable clave", "Semanas (estim.)"],
    [["Etapa 0", "Marca y bases", "Kit de marca listo + auditoría de redes", "Sem 1–2"],
     ["Etapa 1", "Definición", "Catálogo, precios, reglas de anticipo, pago y envíos acordados", "Sem 2–3"],
     ["Etapa 2", "Construcción y piloto", "Tienda bilingüe funcionando; tú administras y probamos los 2 recorridos de compra", "Sem 3–6"],
     ["Etapa 3", "Lanzamiento y mejora", "Sitio en vivo; medimos y ajustamos", "Sem 6 en adelante"]],
    [0.13, 0.20, 0.45, 0.22])
para(doc, "", space_after=2)
runs(para(doc), [("Fecha de inicio propuesta:  ", dict(bold=True, size=10)),
                 ("______________     →     Lanzamiento estimado: ~6 semanas después.", dict(size=10, color=INK_SOFT))])
runs(para(doc), [("Cómo medimos el avance (sin humo):  ", dict(bold=True, color=INK, size=10)),
    ("cotizaciones que terminan en pago, minutos de atención por pedido, entregas a tiempo y "
     "reseñas conseguidas. Los resultados de venta se miden después de lanzar, con datos "
     "reales; no prometemos porcentajes por adelantado.", dict(color=INK_SOFT, size=10))])

# ============================================================== 7. QUÉ NECESITAMOS
section_title(doc, "7", "Lo que necesitamos de ti")
para(doc, "Para hacerte una propuesta con números reales —y no inventados—:", color=INK_SOFT, space_after=6)
for i, q in enumerate([
    "¿Cuántas ventas tienes al mes, en promedio, en los últimos meses? (no el mejor mes, el promedio real).",
    "¿Qué parte de tus ventas es producto de catálogo (precio fijo) y qué parte es 100% personalizado?",
    "Precios y, de ser posible, el margen de tus principales productos (juegos de sellos vs. pines).",
    "¿Cómo cobras hoy y cómo te gustaría cobrar (transferencia, tarjeta, Mercado Pago)? ¿Aplicarías anticipo?",
    "¿Quién va a administrar la tienda en el día a día (tú o alguien de tu equipo)?",
], 1):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(f"{i}.  "); _set_font(r, BODY, 10.5, BLUE, bold=True)
    r = p.add_run(q); _set_font(r, BODY, 10.5, INK_SOFT)
para(doc, "Y para la marca y las redes:", color=INK, bold=True, space_before=6, space_after=4, size=10.5)
bullet(doc, "", "¿Tienes tu logotipo en vector (archivos .ai, .svg, .pdf o .eps)?", dot="—", dotcolor=INK_SOFT)
bullet(doc, "", "¿Qué redes tienes activas además de Facebook (Instagram, TikTok, Google)?", dot="—", dotcolor=INK_SOFT)
bullet(doc, "", "Accesos o nombres de usuario para unificar los perfiles.", dot="—", dotcolor=INK_SOFT)

# ============================================================== 8. INVERSIÓN
section_title(doc, "8", "Inversión (con transparencia)")
para(doc, "Separamos los costos para que veas en qué se va cada peso. El precio de Aimarktech "
          "cubre nuestro trabajo; el dominio, el alojamiento, las comisiones de pago y la "
          "publicidad se cotizan y pagan aparte.", color=INK_SOFT, space_after=8)
simple_table(doc,
    ["Concepto", "Detalle"],
    [["Servicios Aimarktech (mensual)", "Acompañamiento, implementación y ajustes. Planes desde $2,500 (Esencial); $3,500 Crecimiento (recomendado)."],
     ["Kit de marca (si hace falta)", "Diseño o adaptación de logo e identidad para web. Se cotiza según lo que ya tengas."],
     ["Desarrollo inicial de la tienda (una vez)", "Se cotiza al terminar la Etapa 1, cuando conozcamos catálogo y alcance real."],
     ["Servicios externos", "Dominio (~$200–400 MXN/año), alojamiento (desde $0) y correo."],
     ["Comisiones de pago (Opción 2)", "≈ 3.6% + $3 MXN + IVA por transacción. Crecen con las ventas; las cobra la pasarela, no Aimarktech."],
     ["Pauta publicitaria (opcional)", "Ajustable mes con mes; se invierte 100% en la plataforma de anuncios."]],
    [0.34, 0.66])
para(doc, "", space_after=2)
runs(para(doc), [("Nota honesta:  ", dict(bold=True, color=RED, size=10)),
    ("no prometemos un retorno garantizado. Cualquier proyección de ventas la haremos contigo, "
     "con tus números reales, después de la Etapa 1.", dict(color=INK_SOFT, size=10))])

# ============================================================== 9. QUÉ GANA
section_title(doc, "9", "Qué gana Metales Infinity")
bullet(doc, "Tiempo.", "Recuperas horas hoy perdidas en el celular (lo medimos antes y después).", dot="⏱", dotcolor=BLUE)
bullet(doc, "Menos meses flojos.", "Aprovechas la capacidad que ya tienes, con una vitrina que trabaja sola.", dot="📈", dotcolor=BLUE)
bullet(doc, "Clientes que regresan.", "Con seguimiento a tus recompras (pines y sellos), no de memoria.", dot="🔁", dotcolor=BLUE)
bullet(doc, "Enfoque.", "Tiempo para crear piezas nuevas y dirigir tu negocio.", dot="🧠", dotcolor=BLUE)
bullet(doc, "Base para crecer.", "Un camino claro hacia México y EE.UU. con una marca reconocible.", dot="🌎", dotcolor=BLUE)

# ============================================================== 10. POR QUÉ AIMARKTECH
section_title(doc, "10", "Por qué Aimarktech")
for label, txt in [
    ("Diagnóstico antes de recomendar.", "Revisamos el negocio completo antes de proponerte publicidad, software o automatizaciones."),
    ("Una visión conectada.", "Integramos procesos, tecnología, IA y marketing en un mismo plan ejecutable."),
    ("Acompañamiento real.", "No entregamos un documento y desaparecemos: apoyamos la implementación y la transferencia a tu equipo."),
    ("Resultados medibles, sin humo.", "Hablamos claro de límites, dependencias y costos. +28 años de experiencia tecnológica."),
]:
    bullet(doc, label, txt)

# ============================================================== 11. SIGUIENTES PASOS
section_title(doc, "11", "Siguientes pasos")
for i, txt in enumerate([
    "Elige con cuál opción arrancamos (1 o 2).",
    "Una llamada corta (20 min) para responder las preguntas y fijar la fecha de inicio.",
    "Iniciamos la Etapa 0 (marca y bases) y te entregamos el calendario con fechas exactas.",
], 1):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(f"{i}.  "); _set_font(r, BODY, 10.5, BLUE, bold=True)
    r = p.add_run(txt); _set_font(r, BODY, 10.5, INK_SOFT)

contact = doc.add_table(rows=1, cols=1); contact.columns[0].width = CONTENT_W
cc = contact.cell(0, 0); cc.width = CONTENT_W
cell_bg(cc, DARK); cell_margins(cc, 240, 240, 300, 300); cc.paragraphs[0].text = ""
para(cc, "José Antonio Aguilar · Aimarktech", size=12.5, color=WHITE, bold=True, font=HEAD, space_after=4)
runs(para(cc, space_after=2), [("WhatsApp:  ", dict(bold=True, color=CYAN)), ("+52 56 3963 7740", dict(color="D7E4F0"))])
runs(para(cc, space_after=2), [("Oficina:  ", dict(bold=True, color=CYAN)), ("55 8967 7250", dict(color="D7E4F0"))])
runs(para(cc, space_after=2), [("Correo:  ", dict(bold=True, color=CYAN)), ("jantonioab@gmail.com", dict(color="D7E4F0"))])
runs(para(cc, space_after=2), [("Web:  ", dict(bold=True, color=CYAN)), ("https://soyaimarktech.com", dict(color="D7E4F0"))])
para(cc, "Base en Valle de Chalco, con alcance en todo México.", size=9.5, color="B9C6D6", italic=True, space_before=4)

para(doc, "\u201cPrimero trabajamos en el empresario, después en el negocio.\u201d",
     size=10.5, color=BLUE, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12)

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Propuesta_Metales_Infinity_Aimarktech.docx")
doc.save(out)
print("OK ->", out)
